#!/usr/bin/env python3
"""Renders docs/UI-SPEC.md into docs/FiberHomeCPEManager-UI-文稿.docx.

The manuscript and the Word file used to be maintained separately, which meant the docx
silently went stale. This script parses the Markdown subset the manuscript actually uses
(headings, paragraphs, bullets, numbered lists, tables, images, fenced code), so there is
one source of truth.

    <bundled-python> script/make_ui_spec_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "UI-SPEC.md"
OUT = ROOT / "docs" / "FiberHomeCPEManager-UI-文稿.docx"

CJK = "Hiragino Sans GB"   # present on macOS and visible to the renderer
LATIN = "Helvetica Neue"
MONO = "Menlo"
def image_width(path):
    """Panel screenshots are tall, menu-bar strips are wide and short.

    A fixed width made tall shots overflow the page and split their caption onto the next
    page, so the width follows the aspect ratio instead.
    """
    with Image.open(path) as image:
        width, height = image.size
    aspect = height / max(width, 1)
    if aspect > 1.2:
        return Cm(11.5)
    if width < 500:
        return Cm(7)
    # Diagnostic zooms (glyph states) are small assets: at full width they dominate a page.
    if width < 900:
        return Cm(9)
    return Cm(14)

BOLD = re.compile(r"\*\*(.+?)\*\*")
CODE = re.compile(r"`([^`]+)`")
IMAGE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET = re.compile(r"^\s*[-*]\s+(.*)$")
NUMBERED = re.compile(r"^\s*(\d+)\.\s+(.*)$")
TABLE_ROW = re.compile(r"^\|(.+)\|\s*$")
RULE = re.compile(r"^-{3,}\s*$")


def set_cjk(style, size=None, bold=None):
    style.font.name = LATIN
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)


def setup(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    set_cjk(normal, size=10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 11)):
        style = doc.styles[name]
        set_cjk(style, size=size, bold=True)
        style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(4)


def add_runs(paragraph, text, size=None, color=None):
    """Adds text honouring **bold** and `code` spans."""
    for piece in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = MONO
            run.font.size = Pt((size or 10.5) - 1)
        else:
            run = paragraph.add_run(piece)
        if size and not (piece.startswith("`") and piece.endswith("`")):
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    return paragraph


def no_row_split(row, header=False):
    """Keeps a row on one page and repeats the header on continuation pages.

    Without this a long table left an empty fragment row at the top of the next page.
    """
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)
    if header:
        header_flag = OxmlElement("w:tblHeader")
        header_flag.set(qn("w:val"), "true")
        properties.append(header_flag)


def add_shading(cell, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, text in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(text.replace("**", ""))
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.name = LATIN
        run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
        add_shading(cell, "F1F3F5")
    no_row_split(table.rows[0], header=True)
    for row in body:
        created = table.add_row()
        no_row_split(created)
        cells = created.cells
        for index, text in enumerate(row[: len(header)]):
            cells[index].text = ""
            paragraph = cells[index].paragraphs[0]
            add_runs(paragraph, text, size=9.5)
            paragraph.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, alt, relative):
    path = ROOT / relative
    if not path.exists():
        print(f"missing image: {relative}")
        return
    doc.add_picture(str(path), width=image_width(path))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Keep the caption on the same page as its image.
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(alt)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6B, 0x70, 0x76)
    run.font.name = LATIN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    caption.paragraph_format.space_after = Pt(10)


def add_code(doc, lines):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.2)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run("\n".join(lines))
    run.font.name = MONO
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def render(markdown, out_path=OUT):
    doc = Document()
    setup(doc)

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped or RULE.match(stripped):
            index += 1
            continue

        if stripped.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            add_code(doc, block)
            continue

        heading = HEADING.match(stripped)
        if heading:
            level = len(heading.group(1))
            doc.add_heading(heading.group(2).strip(), level=min(level, 4))
            index += 1
            continue

        image = IMAGE.match(stripped)
        if image:
            add_image(doc, image.group(1), image.group(2))
            index += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                raw = lines[index].strip()
                if not re.match(r"^\|[\s:|-]+\|$", raw):
                    cells = [cell.strip() for cell in raw.strip("|").split("|")]
                    rows.append(cells)
                index += 1
            if rows:
                add_table(doc, rows)
            continue

        bullet = BULLET.match(line)
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            add_runs(paragraph, bullet.group(1))
            index += 1
            continue

        numbered = NUMBERED.match(line)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            add_runs(paragraph, numbered.group(2))
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_runs(paragraph, stripped)
        index += 1

    doc.save(out_path)
    return doc


if __name__ == "__main__":
    import sys

    # Any Markdown in the repo can be rendered:  script.py [source.md] [output.docx]
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE
    output = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else OUT
    document = render(source.read_text(), output)
    pictures = len(document.inline_shapes)
    label = output.relative_to(ROOT) if output.is_relative_to(ROOT) else output
    print(f"wrote {label}  paragraphs={len(document.paragraphs)} images={pictures}")
